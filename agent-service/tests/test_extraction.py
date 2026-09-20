from pathlib import Path

import pytest
from docx import Document

from app.extraction import ExtractionError, extract_text


def test_extracts_utf8_text(tmp_path: Path) -> None:
    source = tmp_path / "runbook.txt"
    source.write_text("Fail over, then verify replication.", encoding="utf-8")

    assert extract_text(source)[0].text == "Fail over, then verify replication."


def test_extracts_docx_paragraphs(tmp_path: Path) -> None:
    source = tmp_path / "runbook.docx"
    document = Document()
    document.add_paragraph("First check")
    document.add_paragraph("Second check")
    document.save(source)

    assert extract_text(source)[0].text == "First check\n\nSecond check"


def test_rejects_empty_text(tmp_path: Path) -> None:
    source = tmp_path / "empty.txt"
    source.write_text("   ", encoding="utf-8")

    with pytest.raises(ExtractionError):
        extract_text(source)
